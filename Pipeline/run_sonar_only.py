"""Re-run Stage 5 (SonarQube) alone against an already-built trial dir.

Useful when Stages 1-4 already succeeded (code generated, built, tested) but Stage 5 failed
for an environment reason (e.g. SONAR_TOKEN wasn't set) — avoids re-invoking the LLM and
rebuilding just to retry the Sonar scan.

Usage:
    python3 run_sonar_only.py <trial_dir> [--update-report]

    <trial_dir>      e.g. Runs/2026-07-24_run1/UserStory-01_kimi-k2.6/run1
    --update-report  also patch the matching row in Results/results_<US>_<run_id>.xlsx
                      in place (Sonar columns + Build Error Summary), instead of just
                      printing the result.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl

from config_loader import load_llm_config, load_sonar_config
from language_adapters.csharp_adapter import CSharpAdapter
from report_generator import COLUMNS
from sonar_runner import run_analysis

REPO_ROOT = Path(__file__).resolve().parent.parent
RESULTS_DIR = REPO_ROOT / "Results"


def sanitize_key_component(text: str) -> str:
    return "".join(c if c.isalnum() or c in "-_." else "-" for c in text)


def parse_trial_dir(trial_dir: Path) -> tuple[str, str, str, int]:
    """Return (run_id, user_story_id, model_id, run_number) from a Runs/<...> path."""
    m = re.match(r"^run(\d+)$", trial_dir.name)
    if not m:
        raise ValueError(f"Expected a 'run<N>' directory, got: {trial_dir.name}")
    run_number = int(m.group(1))

    model_folder = trial_dir.parent.name  # "<user_story_id>_<model_id>"
    user_story_id, _, model_id = model_folder.partition("_")
    if not model_id:
        raise ValueError(f"Expected '<user_story_id>_<model_id>' folder, got: {model_folder}")

    run_id = trial_dir.parent.parent.name
    return run_id, user_story_id, model_id, run_number


def update_report_row(
    run_id: str, user_story_id: str, model_display_name: str, run_number: int, sonar_result
) -> Path | None:
    report_path = RESULTS_DIR / f"results_{user_story_id}_{run_id}.xlsx"
    if not report_path.exists():
        print(f"No existing report at {report_path}, skipping update.")
        return None

    wb = openpyxl.load_workbook(report_path)
    ws = wb.active
    header = [c.value for c in ws[1]]
    col = {name: idx + 1 for idx, name in enumerate(header)}

    target_row = None
    for row in ws.iter_rows(min_row=2):
        values = {header[i]: cell.value for i, cell in enumerate(row)}
        if (
            values.get("Run ID") == run_id
            and values.get("User Story ID") == user_story_id
            and values.get("Model") == model_display_name
            and values.get("Run Number") == run_number
        ):
            target_row = row[0].row
            break

    if target_row is None:
        print("No matching row found in report; skipping update.")
        return None

    error_cell = ws.cell(row=target_row, column=col["Build Error Summary"])
    existing_error = error_cell.value or ""
    # Strip only the prior Stage 5 failure segment, keep any earlier-stage errors intact.
    kept = [
        seg for seg in existing_error.split(" | ") if not seg.startswith("Stage 5 (SonarQube)")
    ]
    if sonar_result.error:
        kept.append(f"Stage 5 (SonarQube) failed: {sonar_result.error}")
    error_cell.value = " | ".join(kept)

    if not sonar_result.error:
        severity_str = ", ".join(f"{v} {k}" for k, v in sonar_result.code_smells_by_severity.items())
        ws.cell(row=target_row, column=col["Code Smells (Total)"], value=sonar_result.code_smells_total)
        ws.cell(row=target_row, column=col["Code Smells - by Severity"], value=severity_str)
        ws.cell(row=target_row, column=col["Reliability Rating"], value=sonar_result.reliability_rating)
        ws.cell(row=target_row, column=col["Bug Count"], value=sonar_result.bug_count)
        ws.cell(row=target_row, column=col["Security Rating"], value=sonar_result.security_rating)
        ws.cell(row=target_row, column=col["Vulnerability Count"], value=sonar_result.vulnerability_count)
        ws.cell(row=target_row, column=col["Cyclomatic Complexity"], value=sonar_result.complexity_total)
        ws.cell(row=target_row, column=col["Maintainability Index"], value=sonar_result.maintainability)

    sonar_link = sonar_result.dashboard_url or sonar_result.project_key
    ws.cell(row=target_row, column=col["Sonar Project Key / Dashboard Link"], value=sonar_link)

    wb.save(report_path)
    return report_path


def main(argv: list[str]) -> int:
    if not (1 <= len(argv) <= 2):
        print("Usage: python3 run_sonar_only.py <trial_dir> [--update-report]", file=sys.stderr)
        return 2

    trial_dir = Path(argv[0]).resolve()
    update_report = "--update-report" in argv[1:]

    if not trial_dir.is_dir():
        print(f"Not a directory: {trial_dir}", file=sys.stderr)
        return 2

    run_id, user_story_id, model_id, run_number = parse_trial_dir(trial_dir)

    src_dir = trial_dir / "src"
    if not src_dir.is_dir():
        print(f"No 'src' dir under {trial_dir}", file=sys.stderr)
        return 2

    csproj_path = next(src_dir.rglob("Implementation.csproj"), None)
    if csproj_path is None:
        print("Implementation.csproj not found under src/", file=sys.stderr)
        return 2

    llm_config = load_llm_config()
    sonar_config = load_sonar_config()
    model_config = next((m for m in llm_config.models if m.id == model_id), None)
    if model_config is None:
        print(f"Unknown model id '{model_id}' in {Path('Config/llm.config')}", file=sys.stderr)
        return 2

    adapter = CSharpAdapter()
    sonar_props = adapter.sonar_properties(src_dir)
    project_key = sanitize_key_component(
        f"{sonar_config.project_key_prefix}-{user_story_id}-{model_id}-run{run_number}-{run_id}"
    )

    print(f"Running Sonar analysis for {trial_dir.relative_to(REPO_ROOT)} (project_key={project_key})...")
    result = run_analysis(sonar_config, sonar_props, project_key, csproj_path)

    if result.error:
        print(f"FAILED: {result.error}")
    else:
        print(f"OK: dashboard={result.dashboard_url}")
        print(f"  code_smells={result.code_smells_total} {result.code_smells_by_severity}")
        print(f"  reliability={result.reliability_rating} bugs={result.bug_count}")
        print(f"  security={result.security_rating} vulnerabilities={result.vulnerability_count}")
        print(f"  complexity={result.complexity_total} maintainability={result.maintainability}")

    if update_report:
        updated_path = update_report_row(
            run_id, user_story_id, model_config.display_name, run_number, result
        )
        if updated_path:
            print(f"Report updated: {updated_path.relative_to(REPO_ROOT)}")

    return 0 if not result.error else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
